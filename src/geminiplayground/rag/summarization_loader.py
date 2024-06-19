import logging
import time
import typing
from pathlib import Path
from typing import Iterator

from google import generativeai as genai
from langchain_core.document_loaders import BaseLoader
from langchain_core.documents import Document
from openpyxl.packaging.manifest import mimetypes
from unstructured.partition.auto import partition
from google.generativeai import GenerativeModel
from .caching import cache

summarization_system_instruction = """
 you are an assistant tasked with summarizing files for retrieval. \
 These summaries will be embedded and used to retrieve the original file. Therefore, \
 make sure you detailed describe the content of the file. \
Only describe what you see, hear or read in the file. 
 """


class SummarizationLoader(BaseLoader):

    def __init__(self, model: typing.Union[str, GenerativeModel], *docs: typing.List[str]):
        self._docs = docs
        if isinstance(model, str):
            model = genai.GenerativeModel(model, system_instruction=summarization_system_instruction)
        self._model = model

    # @cache_func_results
    def _get_pdf_parts(self, file_path: typing.Union[str, Path]) -> typing.List[Document]:
        """
        extract the parts of a PDF file.
        """
        file_path = Path(file_path)
        images_folder = file_path.parent / file_path.stem
        images_folder.mkdir(exist_ok=True)
        elements = partition(
            str(file_path),
            strategy="hi_res",
            extract_images_in_pdf=True,
            extract_image_block_types=["Image", "Table"],  # optional
            # can be used to convert them into base64 format
            extract_image_block_to_payload=False,
            extract_image_block_output_dir=str(images_folder),
        )
        return elements

    def _get_pdf_docs(self, file_path: typing.Union[str, Path]) -> Iterator[Document]:
        """
        Process a PDF file and return documents.
        """
        elements = self._get_pdf_parts(file_path)
        for element in elements:
            page_number = element.metadata.page_number
            if element.category == "Image":
                image_path = element.metadata.image_path
                image_summary = f"""
                This is the summarization of the image in the file: {file_path} at page: {page_number}: \n
                {self.summarize_file(image_path)} 
                """
                yield Document(
                    page_content=image_summary,
                    metadata={
                        "file": str(file_path),
                        "image_file": image_path,
                        "page_number": page_number,
                        "category": element.category
                    },
                )
            else:
                yield Document(
                    page_content=element.text,
                    metadata={
                        "file": str(file_path),
                        "page_number": page_number,
                        "category": element.category,
                    },
                )

    def _file_to_docs(self, file_path: typing.Union[str, Path]) -> Iterator[Document]:
        """
        convert a file to documents.
        """
        file_path = Path(file_path)
        mime_type = mimetypes.guess_type(file_path)[0]
        if "pdf" in mime_type:
            yield from self._get_pdf_docs(file_path)
        else:
            yield Document(
                page_content=self.summarize_file(file_path),
                metadata={
                    "file": str(file_path),
                    "mime_type": "text/plain",
                },
            )

    def lazy_load(self) -> Iterator[Document]:
        """
        Lazy load documents.
        """
        for doc_path in self._docs:
            yield from self._file_to_docs(doc_path)

    def summarize_file(self, file_path: typing.Union[str, Path]) -> str:
        """
        summarize a file using Gemini.
        """
        summarization_prompt = """Summarize the following file"""
        file_path = Path(file_path)
        if cache.get(file_path):
            logging.info(f"Loading cached file {file_path}")
            uploaded_file = cache[file_path]
        else:
            logging.info(f"Uploading file {file_path}")
            uploaded_file = genai.upload_file(file_path)
            while uploaded_file.state.name == "PROCESSING":
                time.sleep(10)
                uploaded_file = genai.get_file(uploaded_file.name)
            if uploaded_file.state.name == "FAILED":
                raise Exception("File upload failed")

        response = self._model.generate_content([summarization_prompt, uploaded_file])
        summarization = response.text
        logging.info(f"Summarization for file {file_path}: {summarization}")
        return summarization
